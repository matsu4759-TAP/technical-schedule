"""見積書テンプレート(見積書作成.xlsx)を生成するスクリプト。

使い方: python build_quotation.py
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

FONT = "Meiryo UI"
MAX_LINES = 20          # 見積書に載せられる最大行数
MASTER_ROWS = 60        # 項目マスタの入力可能行数
FIRST = 5               # 項目マスタのデータ開始行

thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F4E78")
input_fill = PatternFill("solid", fgColor="FFF2CC")
yen = '"¥"#,##0;-"¥"#,##0;""'


def f(size=10, bold=False, color="000000"):
    return Font(name=FONT, size=size, bold=bold, color=color)


# 例としてのサンプル項目(自社の項目に置き換えてください)
ITEMS = [
    ("基本費用", "現地調査費", "式", 30000, 1),
    ("基本費用", "設計・計画費", "式", 50000, 1),
    ("基本費用", "諸経費", "式", 20000, 1),
    ("作業費", "技術者作業費", "人日", 45000, 2),
    ("作業費", "測定作業費", "回", 25000, 1),
    ("作業費", "立会い費", "回", 15000, 1),
    ("機材費", "測定機器レンタル", "日", 12000, 2),
    ("機材費", "消耗品", "式", 5000, 1),
    ("交通・出張", "交通費", "式", 10000, 1),
    ("交通・出張", "宿泊費", "泊", 12000, 1),
    ("報告", "報告書作成費", "部", 30000, 1),
    ("報告", "報告会実施費", "回", 20000, 1),
]

wb = Workbook()

# ---------------------------------------------------------------- 項目マスタ
m = wb.active
m.title = "項目選択"
m["A1"] = "見積項目の選択"
m["A1"].font = f(14, True)
m["A2"] = ("「選択」列で TRUE(チェック)にした項目が「見積書」シートに自動で反映されます。"
           "黄色のセルを編集してください。")
m["A2"].font = f(9, color="555555")
m["A3"] = "※Microsoft 365 の場合、A列を選択して［挿入］→［チェックボックス］でクリック式のチェックボックスになります。"
m["A3"].font = f(9, color="555555")

headers = ["選択", "区分", "項目名", "単位", "単価(円)", "数量", "金額(円)", "掲載順"]
widths = [8, 14, 28, 8, 12, 8, 14, 8]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = m.cell(row=FIRST - 1, column=i, value=h)
    c.font = f(10, True, "FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box
    m.column_dimensions[c.column_letter].width = w

last = FIRST + MASTER_ROWS - 1
for r in range(FIRST, last + 1):
    idx = r - FIRST
    if idx < len(ITEMS):
        cat, name, unit, price, qty = ITEMS[idx]
        vals = [False, cat, name, unit, price, qty]
    else:
        vals = [False, None, None, None, None, None]
    for col, v in enumerate(vals, start=1):
        c = m.cell(row=r, column=col, value=v)
        c.fill = input_fill
        c.border = box
        c.font = f()
    m.cell(row=r, column=1).alignment = Alignment(horizontal="center")
    m.cell(row=r, column=5).number_format = "#,##0"
    m.cell(row=r, column=7, value=f'=IF(C{r}="","",E{r}*F{r})').number_format = "#,##0"
    m.cell(row=r, column=8, value=f'=IF(AND(A{r}=TRUE,C{r}<>""),COUNTIF(A${FIRST}:A{r},TRUE),"")')
    for col in (7, 8):
        c = m.cell(row=r, column=col)
        c.border = box
        c.font = f(color="555555" if col == 8 else "000000")
    m.cell(row=r, column=8).alignment = Alignment(horizontal="center")

dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
dv.add(f"A{FIRST}:A{last}")
m.add_data_validation(dv)
# 選択された行をハイライト
m.conditional_formatting.add(
    f"B{FIRST}:G{last}",
    FormulaRule(formula=[f"$A{FIRST}=TRUE"], fill=PatternFill("solid", fgColor="DDEBF7")),
)
m.cell(row=last + 2, column=6, value="選択件数").font = f(10, True)
m.cell(row=last + 2, column=7, value=f"=COUNTIF(A{FIRST}:A{last},TRUE)").font = f(10, True)
m.cell(row=last + 2, column=8,
       value=f'=IF(G{last + 2}>{MAX_LINES},"⚠ {MAX_LINES}件まで","")').font = f(9, True, "C00000")
m.freeze_panes = f"A{FIRST}"

# ---------------------------------------------------------------- 見積書
q = wb.create_sheet("見積書")
for col, w in zip("ABCDEFG", [5, 30, 8, 8, 13, 15, 3]):
    q.column_dimensions[col].width = w

q.merge_cells("A1:F1")
q["A1"] = "御 見 積 書"
q["A1"].font = f(20, True)
q["A1"].alignment = Alignment(horizontal="center")

q["E3"], q["E4"], q["E5"] = "見積番号", "見積日", "有効期限"
q["F3"] = "Q-2026-001"
q["F4"] = "=TODAY()"
q["F5"] = "=F4+30"
for r in (3, 4, 5):
    q[f"E{r}"].font = f(9)
    q[f"F{r}"].font = f(10)
    q[f"F{r}"].alignment = Alignment(horizontal="right")
q["F4"].number_format = q["F5"].number_format = "yyyy年m月d日"
q["F3"].fill = input_fill

q.merge_cells("A3:C3")
q["A3"] = "株式会社サンプル商事"
q["A3"].font = f(14, True)
q["A3"].fill = input_fill
q["D3"] = "御中"
q["D3"].font = f(11)
q.merge_cells("A4:C4")
q["A4"] = "件名：測定業務一式"
q["A4"].fill = input_fill
q["A4"].font = f(10)
q["A6"] = "下記の通り御見積申し上げます。"
q["A6"].font = f(10)

# 自社情報
company = ["〇〇株式会社", "〒000-0000 東京都〇〇区〇〇 1-2-3", "TEL 00-0000-0000", "担当：〇〇"]
for i, line in enumerate(company):
    c = q.cell(row=7 + i, column=5, value=line)
    c.font = f(11 if i == 0 else 9, i == 0)
    c.fill = input_fill

q.merge_cells("A9:B9")
q["A9"] = "御見積金額(税込)"
q["A9"].font = f(11, True)
q.merge_cells("C9:D9")
q["C9"] = "=F{}".format(13 + MAX_LINES + 2)
q["C9"].font = f(16, True)
q["C9"].number_format = yen
for col in "ABCD":
    q[f"{col}9"].border = Border(bottom=Side(style="medium"))

HR = 12  # 明細ヘッダ行
for i, h in enumerate(["No.", "項目名", "数量", "単位", "単価", "金額"], start=1):
    c = q.cell(row=HR, column=i, value=h)
    c.font = f(10, True, "FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(horizontal="center")
    c.border = box

src = "項目選択!${}${}:${}${}"
key = src.format("H", FIRST, "H", last)
def pick(col, n):
    rng = src.format(col, FIRST, col, last)
    return f'IFERROR(INDEX({rng},MATCH({n},{key},0)),"")'

for n in range(1, MAX_LINES + 1):
    r = HR + n
    q.cell(row=r, column=1, value=f'=IF(COUNTIF({key},{n})=0,"",{n})')
    q.cell(row=r, column=2, value=f"={pick('C', n)}")
    q.cell(row=r, column=3, value=f"={pick('F', n)}")
    q.cell(row=r, column=4, value=f"={pick('D', n)}")
    q.cell(row=r, column=5, value=f"={pick('E', n)}")
    q.cell(row=r, column=6, value=f'=IF(B{r}="","",C{r}*E{r})')
    for col in range(1, 7):
        c = q.cell(row=r, column=col)
        c.border = box
        c.font = f()
    for col in (1, 3, 4):
        q.cell(row=r, column=col).alignment = Alignment(horizontal="center")
    q.cell(row=r, column=5).number_format = "#,##0"
    q.cell(row=r, column=6).number_format = "#,##0"

end = HR + MAX_LINES
q["H3"] = "消費税率"
q["I3"] = 0.10
q["I3"].number_format = "0%"
q["I3"].fill = input_fill
q["H3"].font = q["I3"].font = f(9)
totals = [("小計", f"=SUM(F{HR + 1}:F{end})"),
          ("消費税", f"=ROUNDDOWN(F{end + 1}*$I$3,0)"),
          ("合計", f"=F{end + 1}+F{end + 2}")]
for i, (label, formula) in enumerate(totals, start=1):
    r = end + i
    q.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    q.cell(row=r, column=4, value=label).alignment = Alignment(horizontal="center")
    q.cell(row=r, column=6, value=formula).number_format = yen
    for col in (4, 5, 6):
        q.cell(row=r, column=col).border = box
        q.cell(row=r, column=col).font = f(10, label == "合計")

nr = end + 5
q.cell(row=nr, column=1, value="備考").font = f(10, True)
q.merge_cells(start_row=nr + 1, start_column=1, end_row=nr + 3, end_column=6)
c = q.cell(row=nr + 1, column=1, value="・お支払条件：月末締め翌月末払い")
c.alignment = Alignment(vertical="top", wrap_text=True)
c.fill = input_fill
c.font = f(9)
for r in range(nr + 1, nr + 4):
    for col in range(1, 7):
        q.cell(row=r, column=col).border = box

q.print_area = f"A1:F{nr + 3}"
q.page_setup.paperSize = q.PAPERSIZE_A4
q.page_setup.fitToWidth = 1
q.page_setup.fitToHeight = 1
q.sheet_properties.pageSetUpPr.fitToPage = True
q.print_options.horizontalCentered = True
q.sheet_view.showGridLines = False

# ---------------------------------------------------------------- 使い方
h = wb.create_sheet("使い方")
h.column_dimensions["A"].width = 100
lines = [
    ("見積書作成ツールの使い方", f(14, True)),
    ("", None),
    ("1. 「項目選択」シートで、見積書に載せたい項目の「選択」を TRUE(チェック)にする", None),
    ("2. 数量を必要に応じて変更する(単価・項目名の追加/変更も黄色セルで可能)", None),
    ("3. 「見積書」シートで宛先・件名・見積番号など黄色のセルを入力する", None),
    ("4. 「見積書」シートを印刷 または PDF で保存して提出する", None),
    ("", None),
    ("■ ルール", f(11, True)),
    ("・黄色のセル = 入力するセル。白いセルは計算式なので編集しないでください。", None),
    (f"・見積書に載る明細は最大 {MAX_LINES} 行です。", None),
    ("・明細は「項目選択」シートの並び順で掲載されます。", None),
    ("", None),
    ("■ チームでの使い方", f(11, True)),
    ("・このファイルを SharePoint / OneDrive / Teams に置き、マスタ(原本)として管理します。", None),
    ("・見積書を作るときは「名前を付けて保存」でコピーしてから編集すると原本が崩れません。", None),
    ("・マクロ不使用のため、Excel for the web(ブラウザ)や Teams 上でもそのまま使えます。", None),
]
for i, (text, font) in enumerate(lines, start=1):
    c = h.cell(row=i, column=1, value=text)
    c.font = font or f(10)

wb.move_sheet("使い方", offset=-2)
wb.active = 1
wb.save("見積書作成.xlsx")
print("saved 見積書作成.xlsx")
